"""
Cliente para FundsForNGOs Premium.

Replica el flujo del notebook original (scraper_final_v2.ipynb):
1. Cifra la contraseña igual que el frontend Angular.
2. Login por API para obtener un JWT.
3. Búsqueda de convocatorias con filtros (paginada).
4. Detalle completo de cada convocatoria.

Toda la lógica de negocio vive aquí; app.py solo se encarga de la interfaz.
"""

from __future__ import annotations

import base64
import json
import time
from datetime import date, datetime
from io import BytesIO
from typing import Any, Callable

import pandas as pd
import requests
from cryptography.hazmat.backends import default_backend
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from cryptography.hazmat.primitives.padding import PKCS7
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# =============================================================================
# Constantes de la API (mismas del notebook)
# =============================================================================

COOKIE_KEY = b"b14ca5898a4e4133bbce2ea2315a1916"
ZERO_IV = b"\x00" * 16

URL_AUTH = "https://core.fundsforngospremium.com/Accounts/authenticate"
URL_BUSQUEDA = "https://search.fundsforngospremium.com/api/AdvanceGrantSearch/GrantSearch"
URL_DETALLE = "https://core.fundsforngospremium.com/api/Grant/GetGrant"
URL_HIGHLIGHTS = "https://core.fundsforngospremium.com/api/Grant/GetGrantHighlights"

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/145.0.0.0 Safari/537.36"
)

REQUEST_TIMEOUT = (10, 30)
MAX_REINTENTOS_API = 2
PAUSA_REINTENTO_API = 2

# Tope de páginas a escanear cuando se filtra por fecha localmente (ver
# buscar_convocatorias). Evita que una búsqueda con un rango de fechas muy
# angosto sobre un catálogo enorme se quede escaneando indefinidamente.
MAX_PAGINAS_ESCANEADAS = 300


def fecha_en_rango(fecha_str: str, desde: date | None, hasta: date | None) -> bool:
    """Compara la fecha de una convocatoria (string de la API) contra un rango.

    Se usa para filtrar localmente porque la API de FundsForNGOs no aplica
    su propio filtro de fechas de publicación (confirmado en el sitio real).
    Si la fecha no se puede interpretar, se conserva la convocatoria en vez
    de descartarla (mejor mostrar de más que perder resultados válidos).
    """
    if not fecha_str:
        return desde is None and hasta is None
    try:
        fecha = datetime.strptime(fecha_str[:10], "%Y-%m-%d").date()
    except ValueError:
        return True
    if desde and fecha < desde:
        return False
    if hasta and fecha > hasta:
        return False
    return True

URL_PAISES_POR_REGION = "https://core.fundsforngospremium.com/api/Country/GetRegionCountries"

# Catálogos completos de FundsForNGOs Premium.
# Fuente: catálogos exportados por el usuario + verificados contra
# https://core.fundsforngospremium.com/api/Country/GetRegionCountries?id=20

PAISES_DISPONIBLES = [
    {"id": 1, "name": "Afghanistan"},
    {"id": 2, "name": "Aland Islands"},
    {"id": 3, "name": "Albania"},
    {"id": 4, "name": "Algeria"},
    {"id": 5, "name": "American Samoa"},
    {"id": 6, "name": "Andorra"},
    {"id": 7, "name": "Angola"},
    {"id": 8, "name": "Anguilla"},
    {"id": 9, "name": "Antarctica"},
    {"id": 10, "name": "Antigua and Barbuda"},
    {"id": 11, "name": "Argentina"},
    {"id": 12, "name": "Armenia"},
    {"id": 13, "name": "Aruba"},
    {"id": 14, "name": "Australia"},
    {"id": 15, "name": "Austria"},
    {"id": 16, "name": "Azerbaijan"},
    {"id": 17, "name": "Bahamas"},
    {"id": 18, "name": "Bahrain"},
    {"id": 19, "name": "Bangladesh"},
    {"id": 20, "name": "Barbados"},
    {"id": 21, "name": "Belarus"},
    {"id": 22, "name": "Belgium"},
    {"id": 23, "name": "Belize"},
    {"id": 24, "name": "Benin"},
    {"id": 25, "name": "Bermuda"},
    {"id": 26, "name": "Bhutan"},
    {"id": 27, "name": "Bolivia"},
    {"id": 251, "name": "Bonaire Sint Eustatius and Saba"},
    {"id": 28, "name": "Bosnia And Herzegovina"},
    {"id": 29, "name": "Botswana"},
    {"id": 30, "name": "Bouvet Island"},
    {"id": 31, "name": "Brazil"},
    {"id": 32, "name": "British Indian Ocean Territory"},
    {"id": 237, "name": "British Virgin Islands"},
    {"id": 33, "name": "Brunei"},
    {"id": 34, "name": "Bulgaria"},
    {"id": 35, "name": "Burkina Faso"},
    {"id": 36, "name": "Burundi"},
    {"id": 37, "name": "Cambodia"},
    {"id": 38, "name": "Cameroon"},
    {"id": 39, "name": "Canada"},
    {"id": 40, "name": "Cape Verde"},
    {"id": 41, "name": "Cayman Islands"},
    {"id": 42, "name": "Central African Republic"},
    {"id": 43, "name": "Chad"},
    {"id": 44, "name": "Chile"},
    {"id": 45, "name": "China"},
    {"id": 46, "name": "Christmas Island"},
    {"id": 47, "name": "Cocos (Keeling) Islands"},
    {"id": 48, "name": "Colombia"},
    {"id": 49, "name": "Comoros"},
    {"id": 50, "name": "Congo (Brazzaville)"},
    {"id": 51, "name": "Congo DR"},
    {"id": 52, "name": "Cook Islands"},
    {"id": 53, "name": "Costa Rica"},
    {"id": 54, "name": "Cote DIvoire (Ivory Coast)"},
    {"id": 55, "name": "Croatia"},
    {"id": 56, "name": "Cuba"},
    {"id": 250, "name": "CuraÃ§ao"},
    {"id": 57, "name": "Cyprus"},
    {"id": 58, "name": "Czech Republic"},
    {"id": 59, "name": "Denmark"},
    {"id": 60, "name": "Djibouti"},
    {"id": 61, "name": "Dominica"},
    {"id": 62, "name": "Dominican Republic"},
    {"id": 216, "name": "East Timor (Timor-Leste)"},
    {"id": 63, "name": "Ecuador"},
    {"id": 64, "name": "Egypt"},
    {"id": 65, "name": "El Salvador"},
    {"id": 66, "name": "Equatorial Guinea"},
    {"id": 67, "name": "Eritrea"},
    {"id": 68, "name": "Estonia"},
    {"id": 208, "name": "Eswatini"},
    {"id": 69, "name": "Ethiopia"},
    {"id": 70, "name": "Falkland Islands (Malvinas)"},
    {"id": 71, "name": "Faroe Islands"},
    {"id": 72, "name": "Fiji"},
    {"id": 73, "name": "Finland"},
    {"id": 74, "name": "France"},
    {"id": 75, "name": "French Guiana"},
    {"id": 76, "name": "French Polynesia"},
    {"id": 77, "name": "French Southern Territories"},
    {"id": 78, "name": "Gabon"},
    {"id": 79, "name": "Gambia"},
    {"id": 80, "name": "Georgia"},
    {"id": 81, "name": "Germany"},
    {"id": 82, "name": "Ghana"},
    {"id": 83, "name": "Gibraltar"},
    {"id": 84, "name": "Greece"},
    {"id": 85, "name": "Greenland"},
    {"id": 86, "name": "Grenada"},
    {"id": 87, "name": "Guadeloupe"},
    {"id": 88, "name": "Guam"},
    {"id": 89, "name": "Guatemala"},
    {"id": 90, "name": "Guernsey"},
    {"id": 91, "name": "Guinea"},
    {"id": 92, "name": "Guinea-Bissau"},
    {"id": 93, "name": "Guyana"},
    {"id": 94, "name": "Haiti"},
    {"id": 95, "name": "Heard Island And Mcdonald Islands"},
    {"id": 96, "name": "Holy See (Vatican City State)"},
    {"id": 97, "name": "Honduras"},
    {"id": 98, "name": "Hong Kong SAR"},
    {"id": 99, "name": "Hungary"},
    {"id": 100, "name": "Iceland"},
    {"id": 101, "name": "India"},
    {"id": 102, "name": "Indonesia"},
    {"id": 103, "name": "Iran"},
    {"id": 104, "name": "Iraq"},
    {"id": 105, "name": "Ireland"},
    {"id": 106, "name": "Isle Of Man"},
    {"id": 107, "name": "Israel"},
    {"id": 108, "name": "Italy"},
    {"id": 109, "name": "Jamaica"},
    {"id": 110, "name": "Japan"},
    {"id": 111, "name": "Jersey"},
    {"id": 112, "name": "Jordan"},
    {"id": 113, "name": "Kazakhstan"},
    {"id": 114, "name": "Kenya"},
    {"id": 115, "name": "Kiribati"},
    {"id": 247, "name": "Kosovo"},
    {"id": 118, "name": "Kuwait"},
    {"id": 119, "name": "Kyrgyzstan"},
    {"id": 120, "name": "Laos"},
    {"id": 121, "name": "Latvia"},
    {"id": 122, "name": "Lebanon"},
    {"id": 123, "name": "Lesotho"},
    {"id": 124, "name": "Liberia"},
    {"id": 125, "name": "Libya"},
    {"id": 126, "name": "Liechtenstein"},
    {"id": 127, "name": "Lithuania"},
    {"id": 128, "name": "Luxembourg"},
    {"id": 129, "name": "Macau"},
    {"id": 130, "name": "Macedonia"},
    {"id": 131, "name": "Madagascar"},
    {"id": 132, "name": "Malawi"},
    {"id": 133, "name": "Malaysia"},
    {"id": 134, "name": "Maldives"},
    {"id": 135, "name": "Mali"},
    {"id": 136, "name": "Malta"},
    {"id": 137, "name": "Marshall Islands"},
    {"id": 138, "name": "Martinique"},
    {"id": 139, "name": "Mauritania"},
    {"id": 140, "name": "Mauritius"},
    {"id": 141, "name": "Mayotte"},
    {"id": 142, "name": "Mexico"},
    {"id": 143, "name": "Micronesia Federated States Of"},
    {"id": 144, "name": "Moldova Republic Of"},
    {"id": 145, "name": "Monaco"},
    {"id": 146, "name": "Mongolia"},
    {"id": 244, "name": "Montenegro"},
    {"id": 147, "name": "Montserrat"},
    {"id": 148, "name": "Morocco"},
    {"id": 149, "name": "Mozambique"},
    {"id": 150, "name": "Myanmar"},
    {"id": 151, "name": "Namibia"},
    {"id": 152, "name": "Nauru"},
    {"id": 153, "name": "Nepal"},
    {"id": 154, "name": "Netherlands"},
    {"id": 155, "name": "Netherlands Antilles"},
    {"id": 156, "name": "New Caledonia"},
    {"id": 157, "name": "New Zealand"},
    {"id": 158, "name": "Nicaragua"},
    {"id": 159, "name": "Niger"},
    {"id": 160, "name": "Nigeria"},
    {"id": 161, "name": "Niue"},
    {"id": 162, "name": "Norfolk Island"},
    {"id": 117, "name": "North Korea"},
    {"id": 163, "name": "Northern Mariana Islands"},
    {"id": 164, "name": "Norway"},
    {"id": 165, "name": "Oman"},
    {"id": 166, "name": "Pakistan"},
    {"id": 167, "name": "Palau"},
    {"id": 168, "name": "Palestinian Territories"},
    {"id": 169, "name": "Panama"},
    {"id": 170, "name": "Papua New Guinea"},
    {"id": 171, "name": "Paraguay"},
    {"id": 172, "name": "Peru"},
    {"id": 173, "name": "Philippines"},
    {"id": 174, "name": "Pitcairn"},
    {"id": 175, "name": "Poland"},
    {"id": 176, "name": "Portugal"},
    {"id": 177, "name": "Puerto Rico"},
    {"id": 178, "name": "Qatar"},
    {"id": 179, "name": "Reunion"},
    {"id": 180, "name": "Romania"},
    {"id": 181, "name": "Russia"},
    {"id": 182, "name": "Rwanda"},
    {"id": 245, "name": "Saint BarthÃ©lemy"},
    {"id": 183, "name": "Saint Helena"},
    {"id": 184, "name": "Saint Kitts And Nevis"},
    {"id": 185, "name": "Saint Lucia"},
    {"id": 252, "name": "Saint Martin (French Part)"},
    {"id": 186, "name": "Saint Pierre And Miquelon"},
    {"id": 187, "name": "Saint Vincent And The Grenadines"},
    {"id": 188, "name": "Samoa"},
    {"id": 189, "name": "San Marino"},
    {"id": 190, "name": "Sao Tome And Principe"},
    {"id": 191, "name": "Saudi Arabia"},
    {"id": 192, "name": "Senegal"},
    {"id": 246, "name": "Serbia"},
    {"id": 194, "name": "Seychelles"},
    {"id": 195, "name": "Sierra Leone"},
    {"id": 196, "name": "Singapore"},
    {"id": 253, "name": "Sint Maarten (Dutch Part)"},
    {"id": 197, "name": "Slovak Republic"},
    {"id": 198, "name": "Slovenia"},
    {"id": 199, "name": "Solomon Islands"},
    {"id": 200, "name": "Somalia"},
    {"id": 201, "name": "South Africa"},
    {"id": 202, "name": "South Georgia And The South Sandwich Islands"},
    {"id": 116, "name": "South Korea"},
    {"id": 248, "name": "South Sudan"},
    {"id": 203, "name": "Spain"},
    {"id": 204, "name": "Sri Lanka"},
    {"id": 205, "name": "Sudan"},
    {"id": 206, "name": "Suriname"},
    {"id": 207, "name": "Svalbard And Jan Mayen"},
    {"id": 209, "name": "Sweden"},
    {"id": 210, "name": "Switzerland"},
    {"id": 211, "name": "Syria"},
    {"id": 212, "name": "Taiwan"},
    {"id": 213, "name": "Tajikistan"},
    {"id": 214, "name": "Tanzania"},
    {"id": 215, "name": "Thailand"},
    {"id": 217, "name": "Togo"},
    {"id": 218, "name": "Tokelau"},
    {"id": 219, "name": "Tonga"},
    {"id": 220, "name": "Trinidad And Tobago"},
    {"id": 221, "name": "Tunisia"},
    {"id": 222, "name": "Turkey"},
    {"id": 223, "name": "Turkmenistan"},
    {"id": 224, "name": "Turks And Caicos Islands"},
    {"id": 225, "name": "Tuvalu"},
    {"id": 226, "name": "Uganda"},
    {"id": 227, "name": "Ukraine"},
    {"id": 228, "name": "United Arab Emirates"},
    {"id": 229, "name": "United Kingdom"},
    {"id": 230, "name": "United States"},
    {"id": 231, "name": "United States Minor Outlying Islands"},
    {"id": 238, "name": "United States Virgin Islands"},
    {"id": 232, "name": "Uruguay"},
    {"id": 233, "name": "Uzbekistan"},
    {"id": 234, "name": "Vanuatu"},
    {"id": 235, "name": "Venezuela"},
    {"id": 236, "name": "Vietnam"},
    {"id": 239, "name": "Wallis And Futuna"},
    {"id": 240, "name": "Western Sahara"},
    {"id": 241, "name": "Yemen"},
    {"id": 242, "name": "Zambia"},
    {"id": 243, "name": "Zimbabwe"},
]

ISSUES_DISPONIBLES = [
    {"id": 32, "name": "Accidents & Traffic Safety"},
    {"id": 59, "name": "Accounts and Finance"},
    {"id": 3, "name": "Agriculture Food & Nutrition"},
    {"id": 27, "name": "Animals & Wildlife"},
    {"id": 9, "name": "Arts & Culture"},
    {"id": 33, "name": "Business & Industry"},
    {"id": 57, "name": "COVID-19"},
    {"id": 7, "name": "Children"},
    {"id": 8, "name": "Civil Society Development"},
    {"id": 28, "name": "Community Development"},
    {"id": 60, "name": "Construction"},
    {"id": 10, "name": "Democracy & Good Governance"},
    {"id": 12, "name": "Disability"},
    {"id": 26, "name": "Economic Development"},
    {"id": 1, "name": "Education"},
    {"id": 29, "name": "Employment & Labor"},
    {"id": 5, "name": "Environment"},
    {"id": 61, "name": "Equipment"},
    {"id": 52, "name": "Family"},
    {"id": 15, "name": "HIV/AIDS"},
    {"id": 2, "name": "Health"},
    {"id": 30, "name": "Housing & Shelter"},
    {"id": 16, "name": "Human Rights"},
    {"id": 54, "name": "Human Service"},
    {"id": 13, "name": "Humanitarian Relief"},
    {"id": 34, "name": "Immigration"},
    {"id": 18, "name": "Indigenous Communities"},
    {"id": 70, "name": "Individuals"},
    {"id": 17, "name": "Information Technology"},
    {"id": 62, "name": "Innovation"},
    {"id": 35, "name": "International Affairs"},
    {"id": 36, "name": "LGBTQ"},
    {"id": 63, "name": "Leadership"},
    {"id": 19, "name": "Livelihood"},
    {"id": 64, "name": "Logistics"},
    {"id": 65, "name": "Manufacturing"},
    {"id": 66, "name": "Marketing"},
    {"id": 56, "name": "Media"},
    {"id": 31, "name": "Narcotics Drugs & Crime"},
    {"id": 67, "name": "Neighborhood"},
    {"id": 43, "name": "Old Age Care"},
    {"id": 22, "name": "Peace & Conflict Resolution"},
    {"id": 44, "name": "Poverty Alleviation"},
    {"id": 37, "name": "Public Affairs"},
    {"id": 58, "name": "Refugee & Asylum Seekers"},
    {"id": 23, "name": "Religion"},
    {"id": 68, "name": "Research"},
    {"id": 38, "name": "Science"},
    {"id": 39, "name": "Social Sciences"},
    {"id": 53, "name": "Social Service"},
    {"id": 40, "name": "Sports & Recreation"},
    {"id": 55, "name": "Sustainable Development"},
    {"id": 41, "name": "Tourism & Travel"},
    {"id": 69, "name": "Transportation"},
    {"id": 42, "name": "Volunteerism"},
    {"id": 24, "name": "Water & Sanitation"},
    {"id": 14, "name": "Women & Gender"},
    {"id": 25, "name": "Youth & Adolescents"},
]

REGIONES_DISPONIBLES = [
    {"id": 2, "name": "Africa"},
    {"id": 20, "name": "All Countries"},
    {"id": 4, "name": "Americas"},
    {"id": 1, "name": "Asia"},
    {"id": 9, "name": "Asia-Pacific"},
    {"id": 17, "name": "Caribbean"},
    {"id": 32, "name": "Central America"},
    {"id": 33, "name": "Central Asia"},
    {"id": 24, "name": "Central Europe"},
    {"id": 27, "name": "Country Specific"},
    {"id": 16, "name": "Developed Countries"},
    {"id": 19, "name": "Developing Countries"},
    {"id": 34, "name": "East Asia"},
    {"id": 10, "name": "Eastern Africa"},
    {"id": 22, "name": "Eastern Europe"},
    {"id": 26, "name": "Eurasia"},
    {"id": 15, "name": "Europe"},
    {"id": 31, "name": "Global South"},
    {"id": 6, "name": "Latin America"},
    {"id": 18, "name": "Latin America and the Caribbean"},
    {"id": 7, "name": "Mesoamerica"},
    {"id": 11, "name": "Middle Africa"},
    {"id": 8, "name": "Middle East"},
    {"id": 12, "name": "North Africa"},
    {"id": 25, "name": "North America"},
    {"id": 30, "name": "Oceania"},
    {"id": 29, "name": "Pacific"},
    {"id": 28, "name": "South America"},
    {"id": 13, "name": "South Asia"},
    {"id": 3, "name": "Southeast Asia"},
    {"id": 14, "name": "Southern Africa"},
    {"id": 21, "name": "Sub-Saharan Africa"},
    {"id": 5, "name": "Western Africa"},
    {"id": 23, "name": "Western Europe"},
]

TAMANOS_GRANT_DISPONIBLES = [
    {"id": 0, "name": "Less than $1000"},
    {"id": 1, "name": "$1000 to $10,000"},
    {"id": 2, "name": "$10,000 to $100,000"},
    {"id": 3, "name": "$100,000 to $500,000"},
    {"id": 4, "name": "$500,000 to $1 million"},
    {"id": 5, "name": "More than $1 million"},
]

TIPOS_GRANT_DISPONIBLES = [
    {"id": 21, "name": "Awards, Prizes and Challenges"},
    {"id": 26, "name": "Events"},
    {"id": 11, "name": "Fellowship"},
    {"id": 22, "name": "Grant"},
    {"id": 23, "name": "Scholarship"},
    {"id": 28, "name": "Training"},
]




class FFNAuthError(RuntimeError):
    """Login rechazado o credenciales inválidas."""


class FFNClient:
    """Mantiene la sesión (token) y expone búsqueda + detalle."""

    def __init__(self, email: str, password: str):
        self.email = email
        self.password = password
        self.http = requests.Session()
        self.token = ""
        self.member_id = ""
        self.login_response: dict[str, Any] = {}

    # ---------------------------------------------------------------
    # Autenticación
    # ---------------------------------------------------------------
    @staticmethod
    def _cifrar_password(password: str) -> str:
        padder = PKCS7(128).padder()
        padded = padder.update(password.encode("utf-8")) + padder.finalize()

        cipher = Cipher(
            algorithms.AES(COOKIE_KEY), modes.CBC(ZERO_IV), backend=default_backend()
        )
        encryptor = cipher.encryptor()
        encrypted = encryptor.update(padded) + encryptor.finalize()

        token = base64.b64encode(encrypted).decode("ascii")
        return token.replace("=", "|").replace("/", "!").replace("+", "-")

    @staticmethod
    def _decodificar_jwt_payload(token: str) -> dict[str, Any]:
        try:
            payload_b64 = token.split(".")[1]
            payload_b64 += "=" * (-len(payload_b64) % 4)
            return json.loads(base64.urlsafe_b64decode(payload_b64).decode("utf-8"))
        except Exception:
            return {}

    def login(self) -> None:
        payload = {
            "email": self.email,
            "password": self._cifrar_password(self.password),
            "RefURL": "",
            "userIP": "",
        }
        headers = {
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json",
            "Origin": "https://login.fundsforngospremium.com",
            "Referer": "https://login.fundsforngospremium.com/",
            "User-Agent": USER_AGENT,
        }

        try:
            respuesta = self.http.post(
                URL_AUTH, json=payload, headers=headers, timeout=REQUEST_TIMEOUT
            )
        except requests.RequestException as exc:
            raise FFNAuthError(f"No se pudo contactar el servidor de login: {exc}") from exc

        if respuesta.status_code != 200:
            raise FFNAuthError(f"Login falló con HTTP {respuesta.status_code}")

        datos = respuesta.json()
        if datos.get("status") != 1:
            raise FFNAuthError(datos.get("message", "Login rechazado por el servidor"))

        token = (
            datos.get("tokenResponse", {}).get("jwtToken")
            or datos.get("jwtToken")
            or datos.get("token")
            or ""
        )
        if not token:
            raise FFNAuthError("Login exitoso pero la respuesta no incluyó JWT")

        self.token = token
        self.login_response = datos
        self.member_id = self._extraer_member_id()

    def _extraer_member_id(self) -> str:
        session = (self.login_response or {}).get("sessionVariable", {})
        for clave in ["memberId", "memberIndex", "MemberId", "member_id", "userId"]:
            valor = session.get(clave)
            if valor:
                return str(valor)

        payload = self._decodificar_jwt_payload(self.token)
        for clave in ["MemberId", "memberId", "member_id", "id", "sub", "userId"]:
            valor = payload.get(clave)
            if valor:
                return str(valor)
        return ""

    def _renovar_token(self) -> bool:
        try:
            self.login()
            return True
        except FFNAuthError:
            return False

    def _headers(self) -> dict[str, str]:
        return {
            "Accept": "application/json, text/plain, */*",
            "Authorization": f"Bearer {self.token}",
            "Content-Type": "application/json",
            "Origin": "https://grants.fundsforngospremium.com",
            "Referer": "https://grants.fundsforngospremium.com/",
            "User-Agent": USER_AGENT,
        }

    def _request(self, metodo: str, url: str, **kwargs: Any) -> requests.Response:
        ultimo_error = None
        for intento in range(1, MAX_REINTENTOS_API + 1):
            try:
                return self.http.request(
                    metodo, url, headers=self._headers(), timeout=REQUEST_TIMEOUT, **kwargs
                )
            except requests.RequestException as exc:
                ultimo_error = exc
                if intento < MAX_REINTENTOS_API:
                    time.sleep(PAUSA_REINTENTO_API)
        raise RuntimeError(f"No se pudo consultar {url}: {ultimo_error}")

    # ---------------------------------------------------------------
    # Búsqueda
    # ---------------------------------------------------------------
    @staticmethod
    def _formato_fecha_api(valor: Any) -> str:
        """La API espera fechas SIN ceros a la izquierda: '2026-9-1', no
        '2026-09-01' (confirmado capturando el payload real del sitio)."""
        if not valor:
            return ""
        if isinstance(valor, str):
            return valor
        return f"{valor.year}-{valor.month}-{valor.day}"

    def _payload_busqueda(self, filtros: dict[str, Any], pagina: int) -> dict[str, Any]:
        primera_region = filtros["regions"][0]["id"] if filtros["regions"] else 0

        # "fromDate"/"toDate" (fecha de PUBLICACIÓN) nunca se envían al
        # servidor: se confirmó que su propio filtro de fecha de publicación
        # no funciona del lado de FFN, así que ese rango se aplica solo
        # localmente (ver fecha_en_rango / buscar_convocatorias).
        #
        # La fecha de CIERRE (deadline) sí funciona en el servidor, pero
        # usa el mismo par de campos fromDate/toDate, activado con
        # filterBy="DD" (confirmado capturando el payload real del sitio).
        deadline_activo = bool(filtros.get("deadlineActivo"))
        if deadline_activo:
            filter_by = "DD"
            fecha_desde_api = self._formato_fecha_api(filtros.get("deadlineDesde"))
            fecha_hasta_api = self._formato_fecha_api(filtros.get("deadlineHasta"))
        else:
            filter_by = "-"
            fecha_desde_api = ""
            fecha_hasta_api = ""

        return {
            "pageIndex": pagina,
            "pageSize": 10,
            "platform": "NEW",
            "viewType": "GLOBAL",
            "regionID": primera_region,
            "regions": filtros["regions"],
            "countries": filtros["countries"],
            "countryFlag": len(filtros["countries"]) > 0,
            "issues": filtros["issues"],
            "grantSizes": filtros["grantSizes"],
            "grantTypes": filtros["grantTypes"],
            "donors": [],
            "filterBy": filter_by,
            "SearchText": filtros["SearchText"],
            "fromDate": fecha_desde_api,
            "toDate": fecha_hasta_api,
            "grantSizeFlag": len(filtros["grantSizes"]) > 0,
            "grantTypeFlag": len(filtros["grantTypes"]) > 0,
            "issueFlag": len(filtros["issues"]) > 0,
            "includeExpired": filtros["includeExpired"],
            "deadlineFreeGrants": False,
            "memberId": self.member_id or "",
            "userIp": "",
        }

    def _post_busqueda(self, payload: dict[str, Any]) -> requests.Response:
        respuesta = self._request("POST", URL_BUSQUEDA, json=payload)
        if respuesta.status_code == 401 and self._renovar_token():
            respuesta = self._request("POST", URL_BUSQUEDA, json=payload)
        return respuesta

    def buscar_pagina(self, filtros: dict[str, Any], pagina: int = 1) -> tuple[list[dict], int]:
        """Trae una sola página. Útil para el conteo/preview rápido."""
        payload = self._payload_busqueda(filtros, pagina)
        respuesta = self._post_busqueda(payload)
        if respuesta.status_code != 200:
            raise RuntimeError(f"Error HTTP {respuesta.status_code} consultando resultados")
        datos = respuesta.json()
        return datos.get("result", []), datos.get("recordCount", 0)

    def buscar_convocatorias(
        self,
        filtros: dict[str, Any],
        limite: int | None = None,
        pausa_entre_paginas: float = 0.5,
        on_progreso: Callable[[int, int, int, int], None] | None = None,
        max_paginas: int = MAX_PAGINAS_ESCANEADAS,
    ) -> list[dict[str, Any]]:
        """Busca convocatorias, paginando.

        IMPORTANTE: el filtro de fechas de publicación NO funciona del lado
        del servidor de FundsForNGOs (verificado directamente en su sitio,
        no solo en esta app). Por eso, si `filtros['fromDate']` o
        `filtros['toDate']` vienen con una fecha, aquí se descarta
        localmente cualquier resultado fuera de ese rango, revisando todas
        las páginas necesarias (no solo hasta llegar al `limite`), hasta un
        máximo de `max_paginas` por seguridad.
        """
        convocatorias: list[dict[str, Any]] = []
        pagina = 1
        total = None
        desde = filtros.get("fromDate")
        hasta = filtros.get("toDate")
        filtrar_fecha_local = isinstance(desde, date) or isinstance(hasta, date)

        while True:
            resultados, total_actual = self.buscar_pagina(filtros, pagina)
            if total is None:
                total = total_actual

            if not resultados:
                break

            if filtrar_fecha_local:
                resultados = [
                    r for r in resultados if fecha_en_rango(r.get("lastUpdated", ""), desde, hasta)
                ]

            espacio = None if limite is None else max(limite - len(convocatorias), 0)
            convocatorias.extend(resultados if espacio is None else resultados[:espacio])

            pagina_maxima = min(-(-total // 10), max_paginas) if total else pagina
            if on_progreso:
                on_progreso(len(convocatorias), total or 0, pagina, pagina_maxima)

            if limite is not None and len(convocatorias) >= limite:
                break
            if not filtrar_fecha_local and total is not None and len(convocatorias) >= total:
                break
            if pagina >= pagina_maxima:
                break

            pagina += 1
            time.sleep(pausa_entre_paginas)

        return convocatorias

    # ---------------------------------------------------------------
    # Detalle
    # ---------------------------------------------------------------
    def _get_detalle(self, params: dict[str, str]) -> requests.Response:
        respuesta = self._request("GET", URL_DETALLE, params=params)
        if respuesta.status_code == 401 and self._renovar_token():
            params["MemberId"] = self.member_id or ""
            respuesta = self._request("GET", URL_DETALLE, params=params)
        return respuesta

    def _get_highlights(self, params: dict[str, str]) -> requests.Response:
        respuesta = self._request("GET", URL_HIGHLIGHTS, params=params)
        if respuesta.status_code == 401 and self._renovar_token():
            params["MemberId"] = self.member_id or ""
            respuesta = self._request("GET", URL_HIGHLIGHTS, params=params)
        return respuesta

    @staticmethod
    def _primer_valor(datos: dict[str, Any], claves: list[str]) -> Any:
        for clave in claves:
            valor = datos.get(clave)
            if valor not in (None, "", []):
                return valor
        return ""

    @staticmethod
    def _nombres_desde_lista(valor: Any) -> str:
        if not valor:
            return ""
        if isinstance(valor, str):
            return valor
        if isinstance(valor, dict):
            valor = [valor]
        if not isinstance(valor, list):
            return ""

        nombres = []
        for item in valor:
            if isinstance(item, str):
                nombre = item
            elif isinstance(item, dict):
                nombre = FFNClient._primer_valor(
                    item,
                    ["name", "title", "country", "countryName", "issue", "issueName", "focusAreaName"],
                )
            else:
                nombre = ""
            if nombre:
                nombres.append(str(nombre))
        return ", ".join(nombres)

    @staticmethod
    def _normalizar_respuesta_detalle(datos: dict[str, Any]) -> dict[str, Any]:
        resultado = datos.get("grant") or datos.get("result") or datos
        if isinstance(resultado, list):
            resultado = resultado[0] if resultado else {}
        if isinstance(resultado, dict) and isinstance(resultado.get("grant"), dict):
            resultado = resultado["grant"]
        return resultado if isinstance(resultado, dict) else {}

    def obtener_detalle(self, slug: str) -> dict[str, str]:
        vacio = {
            "descripcion_completa": "",
            "url_referencia": "",
            "areas_tematicas": "",
            "paises_elegibles": "",
        }
        if not slug:
            return vacio

        try:
            respuesta = self._get_detalle({"url": slug, "MemberId": self.member_id or ""})
        except RuntimeError as exc:
            vacio["descripcion_completa"] = f"Error consultando detalle: {exc}"
            return vacio

        if respuesta.status_code != 200:
            vacio["descripcion_completa"] = f"Error HTTP {respuesta.status_code}"
            return vacio

        datos = respuesta.json()
        grant = self._normalizar_respuesta_detalle(datos)
        highlights: dict[str, Any] = {}
        ref_index = grant.get("id")
        if ref_index:
            try:
                highlights_respuesta = self._get_highlights(
                    {"RefIndex": str(ref_index), "MemberId": self.member_id or ""}
                )
                if highlights_respuesta.status_code == 200:
                    highlights = highlights_respuesta.json()
            except RuntimeError:
                pass

        descripcion = self._primer_valor(
            grant,
            ["text", "fullDescription", "grantDescription", "content", "details", "about", "description", "summary"],
        )
        url_referencia = self._primer_valor(
            grant,
            ["referenceUrl", "referenceURL", "reference_url", "externalUrl", "websiteUrl", "link", "urlReference", "url"],
        )
        areas = self._primer_valor(grant, ["focusAreas", "issueAreas", "issues", "themes", "areas", "grantFocusAreas"])
        if not areas:
            areas = self._primer_valor(highlights, ["focusIssues", "issueAreas", "issues", "themes", "areas"])
        paises = self._primer_valor(
            grant, ["focusCountries", "countries", "eligibleCountries", "grantCountries", "country", "countriesEligible"]
        )
        if not paises:
            paises = self._primer_valor(highlights, ["focusCountries", "focusCounties", "countries", "eligibleCountries"])

        return {
            "descripcion_completa": descripcion or "",
            "url_referencia": url_referencia or "",
            "areas_tematicas": self._nombres_desde_lista(areas),
            "paises_elegibles": self._nombres_desde_lista(paises),
        }


# =============================================================================
# Utilidades compartidas por la app
# =============================================================================

def obtener_paises_por_region(region_id: int) -> list[dict[str, Any]]:
    """Consulta el catálogo de países filtrado por región (endpoint público, sin token).

    Se usa para el filtro en cascada Región -> Países en la interfaz.
    """
    try:
        respuesta = requests.get(
            URL_PAISES_POR_REGION, params={"id": region_id}, timeout=REQUEST_TIMEOUT
        )
        respuesta.raise_for_status()
        datos = respuesta.json()
    except (requests.RequestException, ValueError) as exc:
        raise RuntimeError(f"No se pudo consultar países de la región {region_id}: {exc}") from exc

    resultado: list[dict[str, Any]] = []
    vistos: set[int] = set()
    for pais in datos.get("countries", []):
        cid = pais.get("countryIndex")
        nombre = (pais.get("countryName") or "").strip()
        if cid is None or cid in vistos or not nombre:
            continue
        vistos.add(cid)
        resultado.append({"id": cid, "name": nombre})

    resultado.sort(key=lambda x: x["name"])
    return resultado


def formatear_fecha(fecha_str: str) -> str:
    if not fecha_str:
        return ""
    try:
        return datetime.strptime(fecha_str[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return fecha_str[:10]


def construir_dataframe(
    client: FFNClient,
    listado: list[dict[str, Any]],
    on_progreso: Callable[[int, int], None] | None = None,
    pausa_entre_detalles: float = 0.3,
) -> pd.DataFrame:
    filas = []
    total = len(listado)
    for i, conv in enumerate(listado, start=1):
        slug = conv.get("url", "")
        detalle = client.obtener_detalle(slug)

        filas.append(
            {
                "Titulo": conv.get("title", ""),
                "Categoria": conv.get("category", ""),
                "Donante": conv.get("donorAgency", ""),
                "Monto": conv.get("grantSize", ""),
                "Fecha publicacion": formatear_fecha(conv.get("lastUpdated", "")),
                "Fecha limite": formatear_fecha(conv.get("deadline", "")),
                "Descripcion corta": conv.get("description", ""),
                "Descripcion completa": detalle["descripcion_completa"],
                "URL referencia": detalle["url_referencia"],
                "Areas tematicas": detalle["areas_tematicas"],
                "Paises elegibles": detalle["paises_elegibles"],
                "URL en plataforma": f"https://grants.fundsforngospremium.com/opportunity/op/{slug}",
            }
        )
        if on_progreso:
            on_progreso(i, total)
        time.sleep(pausa_entre_detalles)

    return pd.DataFrame(filas)


def exportar_excel(df: pd.DataFrame) -> bytes:
    """Genera el .xlsx con el mismo formato del notebook, en memoria (sin tocar disco)."""
    buffer = BytesIO()
    df.to_excel(buffer, sheet_name="Convocatorias", index=False, engine="openpyxl")
    buffer.seek(0)

    wb = load_workbook(buffer)
    ws = wb["Convocatorias"]

    fuente_header = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    fondo_header = PatternFill("solid", start_color="1F4E79")
    alin_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alin_datos = Alignment(vertical="top", wrap_text=True)
    fuente_datos = Font(name="Arial", size=9)

    anchos = {
        "Titulo": 50,
        "Categoria": 22,
        "Donante": 32,
        "Monto": 24,
        "Fecha publicacion": 18,
        "Fecha limite": 14,
        "Descripcion corta": 55,
        "Descripcion completa": 80,
        "URL referencia": 55,
        "Areas tematicas": 40,
        "Paises elegibles": 40,
        "URL en plataforma": 65,
    }

    for col_idx, nombre in enumerate(df.columns, start=1):
        celda = ws.cell(row=1, column=col_idx)
        celda.font = fuente_header
        celda.fill = fondo_header
        celda.alignment = alin_header
        ws.column_dimensions[get_column_letter(col_idx)].width = anchos.get(nombre, 20)

    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for celda in fila:
            celda.font = fuente_datos
            celda.alignment = alin_datos

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida.getvalue()


def exportar_excel_generico(df: pd.DataFrame, nombre_hoja: str = "Datos") -> bytes:
    """Genera un .xlsx con encabezado estilizado para cualquier DataFrame.

    Versión simplificada de exportar_excel() (sin anchos de columna
    específicos por nombre) para las tablas del apartado de evaluación CTI.
    """
    hoja = nombre_hoja[:31] or "Datos"
    buffer = BytesIO()
    df.to_excel(buffer, sheet_name=hoja, index=False, engine="openpyxl")
    buffer.seek(0)

    wb = load_workbook(buffer)
    ws = wb[hoja]

    fuente_header = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    fondo_header = PatternFill("solid", start_color="1F4E79")
    alin_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alin_datos = Alignment(vertical="top", wrap_text=True)
    fuente_datos = Font(name="Arial", size=9)

    for col_idx, nombre in enumerate(df.columns, start=1):
        celda = ws.cell(row=1, column=col_idx)
        celda.font = fuente_header
        celda.fill = fondo_header
        celda.alignment = alin_header
        ancho = min(max(len(str(nombre)) + 6, 14), 60)
        ws.column_dimensions[get_column_letter(col_idx)].width = ancho

    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for celda in fila:
            celda.font = fuente_datos
            celda.alignment = alin_datos

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    salida = BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida.getvalue()
